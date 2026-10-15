import csv
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from reportlab.platypus import Image as PdfImage, Paragraph, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
REPORT = ROOT / "report"
PDF_PATH = REPORT / "full_project_report.pdf"
PPT_PATH = REPORT / "project_presentation.pptx"

NAVY = "#142B45"
INK = "#213449"
BLUE = "#2878A5"
CYAN = "#48B8C5"
MINT = "#79C8AE"
PALE = "#EEF5F8"
PALE_MINT = "#EAF6F2"
GOLD = "#F2B84B"
CORAL = "#E97868"
WHITE = "#FFFFFF"
MUTED = "#657789"
LIGHT_LINE = "#D9E4EA"

def read_transcript(name):
    path = OUT / name
    if not path.exists():
        return ""
    text = path.read_text(encoding="utf-8").strip()
    if name == "asr_transcription.txt":
        lines = text.splitlines()
        if len(lines) >= 3 and lines[1].startswith("="):
            return "\n".join(lines[2:]).strip()
    return text


def load_metrics():
    with (OUT / "wer_results.csv").open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def audio_stats():
    import soundfile as sf

    source = sf.info(str(ROOT / "audio" / "stutter_current.mp3"))
    tts = sf.info(str(OUT / "tts_output.wav"))
    return {
        "source_duration": source.duration,
        "source_rate": source.samplerate,
        "source_channels": source.channels,
        "tts_duration": tts.duration,
        "tts_rate": tts.samplerate,
        "tts_channels": tts.channels,
    }


def clean_markup(text):
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


PDF_STYLES = {
    "body": ParagraphStyle(
        "Body", fontName="Helvetica", fontSize=8.6, leading=12,
        textColor=colors.HexColor(INK), spaceAfter=4,
    ),
    "small": ParagraphStyle(
        "Small", fontName="Helvetica", fontSize=7.5, leading=10,
        textColor=colors.HexColor(MUTED),
    ),
    "callout": ParagraphStyle(
        "Callout", fontName="Helvetica", fontSize=8.5, leading=12,
        textColor=colors.HexColor(INK),
    ),
    "table": ParagraphStyle(
        "Table", fontName="Helvetica", fontSize=7.4, leading=9.4,
        textColor=colors.HexColor(INK),
    ),
    "table_head": ParagraphStyle(
        "TableHead", fontName="Helvetica-Bold", fontSize=7.2, leading=9,
        textColor=colors.white,
    ),
}


def pdf_para(c, x, top, width, text, kind="body"):
    para = Paragraph(text, PDF_STYLES[kind])
    _, height = para.wrap(width, 1000)
    para.drawOn(c, x, top - height)
    return top - height


def pdf_plain(c, x, y, text, size=9, color=INK, bold=False):
    c.setFillColor(colors.HexColor(color))
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    c.drawString(x, y, text)


def pdf_page(c, number, title, subtitle):
    width, height = A4
    c.setFillColor(colors.white)
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(NAVY))
    c.rect(0, height - 13 * mm, width, 13 * mm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(CYAN))
    c.rect(0, height - 13 * mm, 7 * mm, 13 * mm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica-Bold", 7)
    c.drawString(18 * mm, height - 8.5 * mm, "STUTTER SPEECH | ASR + TTS")
    c.setFillColor(colors.HexColor(NAVY))
    c.setFont("Helvetica-Bold", 18)
    c.drawString(18 * mm, height - 27 * mm, title)
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica", 8.2)
    c.drawString(18 * mm, height - 33 * mm, subtitle)
    c.setStrokeColor(colors.HexColor(LIGHT_LINE))
    c.line(18 * mm, 13 * mm, width - 18 * mm, 13 * mm)
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica", 7)
    c.drawString(18 * mm, 8 * mm, "Educational demonstration - not a clinical assessment")
    c.drawRightString(width - 18 * mm, 8 * mm, f"{number:02d} / 15")
    return height - 43 * mm


def pdf_section(c, x, top, title, width=175 * mm):
    c.setFillColor(colors.HexColor(BLUE))
    c.roundRect(x, top - 5.5 * mm, 2.2 * mm, 5.5 * mm, 1 * mm, fill=1, stroke=0)
    pdf_plain(c, x + 5 * mm, top - 4.3 * mm, title, size=10.2, color=NAVY, bold=True)
    return top - 8 * mm


def pdf_card(c, x, top, width, height, title, text, fill=PALE):
    c.setFillColor(colors.HexColor(fill))
    c.setStrokeColor(colors.HexColor(LIGHT_LINE))
    c.roundRect(x, top - height, width, height, 7, fill=1, stroke=1)
    pdf_plain(c, x + 8, top - 15, title, size=8.4, color=BLUE, bold=True)
    pdf_para(c, x + 8, top - 22, width - 16, text)
    return top - height


def pdf_image(c, path, x, top, width, max_height):
    image = PdfImage(str(path))
    iw, ih = image.imageWidth, image.imageHeight
    scale = min(width / iw, max_height / ih)
    draw_w, draw_h = iw * scale, ih * scale
    c.drawImage(str(path), x + (width - draw_w) / 2, top - draw_h, draw_w, draw_h, preserveAspectRatio=True)
    return top - draw_h


def pdf_table(c, x, top, widths, rows, header=True):
    cooked = []
    for row_index, row in enumerate(rows):
        style = PDF_STYLES["table_head"] if header and row_index == 0 else PDF_STYLES["table"]
        cooked.append([Paragraph(str(value), style) for value in row])
    table = Table(cooked, colWidths=widths, repeatRows=1 if header else 0)
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(LIGHT_LINE)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(PALE)]),
        ]
    else:
        commands += [("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor(PALE)])]
    table.setStyle(TableStyle(commands))
    _, height = table.wrap(sum(widths), 1000)
    table.drawOn(c, x, top - height)
    return top - height


def draw_pdf_cover(c, stats):
    width, height = A4
    c.setFillColor(colors.HexColor(NAVY))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(BLUE))
    c.circle(width + 3 * mm, height - 25 * mm, 62 * mm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(CYAN))
    c.circle(width - 1 * mm, height - 22 * mm, 41 * mm, fill=1, stroke=0)
    c.setFillColor(colors.HexColor(MINT))
    c.circle(16 * mm, 34 * mm, 28 * mm, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm, height - 33 * mm, "PROJECT REPORT  |  SPEECH TECHNOLOGY")
    c.setFont("Helvetica-Bold", 27)
    c.drawString(20 * mm, height - 75 * mm, "From Voice to")
    c.drawString(20 * mm, height - 89 * mm, "Text and Back")
    c.setFillColor(colors.HexColor("#D8E8F1"))
    c.setFont("Helvetica", 12)
    c.drawString(20 * mm, height - 103 * mm, "ASR, CTC, WER and text-to-speech")
    c.drawString(20 * mm, height - 111 * mm, "analysis of a speech sample with disfluencies")
    c.setStrokeColor(colors.HexColor(CYAN))
    c.setLineWidth(2)
    c.line(20 * mm, height - 124 * mm, 90 * mm, height - 124 * mm)
    stages = ["AUDIO", "FEATURES", "ASR", "TEXT", "TTS"]
    x = 20 * mm
    y = height - 158 * mm
    for i, stage in enumerate(stages):
        c.setFillColor(colors.HexColor([BLUE, CYAN, MINT, BLUE, CORAL][i]))
        c.roundRect(x, y, 30 * mm, 12 * mm, 3 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 7.5)
        c.drawCentredString(x + 15 * mm, y + 4.4 * mm, stage)
        if i < len(stages) - 1:
            c.setFillColor(colors.HexColor("#D8E8F1"))
            c.setFont("Helvetica-Bold", 12)
            c.drawCentredString(x + 34 * mm, y + 4.3 * mm, ">")
        x += 37 * mm
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(20 * mm, height - 184 * mm, "Sample profile")
    c.setFillColor(colors.HexColor("#D8E8F1"))
    c.setFont("Helvetica", 9)
    c.drawString(
        20 * mm,
        height - 192 * mm,
        f"English MP3 | {stats['source_duration']:.2f} seconds | {stats['source_rate']:,} Hz",
    )
    c.drawString(20 * mm, height - 199 * mm, "Wav2Vec2-base-960h + Whisper-base | mono ASR input")
    c.setFillColor(colors.HexColor("#D8E8F1"))
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, 20 * mm, "Educational demonstration - not a clinical assessment")
    c.drawRightString(width - 20 * mm, 20 * mm, "01 / 15")
    c.showPage()


def pdf_pipeline(c, x, top, width):
    labels = [
        ("Speech", "MP3"),
        ("Features", "wave + Mel"),
        ("Recognizer", "2 model paths"),
        ("Evaluation", "CTC / WER"),
        ("Synthesis", "TTS WAV"),
    ]
    gap = 6
    boxw = (width - 4 * gap) / 5
    y = top - 32
    for i, (a, b) in enumerate(labels):
        bx = x + i * (boxw + gap)
        c.setFillColor(colors.HexColor(PALE if i % 2 == 0 else PALE_MINT))
        c.setStrokeColor(colors.HexColor(BLUE if i % 2 == 0 else CYAN))
        c.roundRect(bx, y, boxw, 31, 5, fill=1, stroke=1)
        c.setFillColor(colors.HexColor(NAVY))
        c.setFont("Helvetica-Bold", 7.5)
        c.drawCentredString(bx + boxw / 2, y + 18, a)
        c.setFillColor(colors.HexColor(MUTED))
        c.setFont("Helvetica", 6.6)
        c.drawCentredString(bx + boxw / 2, y + 7, b)
        if i < len(labels) - 1:
            c.setFillColor(colors.HexColor(BLUE))
            c.setFont("Helvetica-Bold", 9)
            c.drawCentredString(bx + boxw + gap / 2, y + 13, ">")
    return y - 8


def pdf_bar_chart(c, x, top, width, rows):
    chart_h = 85
    left = x + 33
    right = x + width - 8
    bottom = top - chart_h
    c.setStrokeColor(colors.HexColor(LIGHT_LINE))
    c.line(left, bottom, right, bottom)
    max_val = 100
    for tick in (0, 20, 40, 60, 80, 100):
        tx = left + (right - left) * tick / max_val
        c.setStrokeColor(colors.HexColor(LIGHT_LINE))
        c.line(tx, bottom, tx, bottom + chart_h - 12)
        c.setFillColor(colors.HexColor(MUTED))
        c.setFont("Helvetica", 6.5)
        c.drawCentredString(tx, bottom - 10, str(tick))
    colors_by_model = [BLUE, CORAL]
    for i, row in enumerate(rows[:2]):
        value = float(row["wer"]) * 100
        yy = top - 22 - i * 30
        barw = (right - left) * value / max_val
        c.setFillColor(colors.HexColor(colors_by_model[i]))
        c.roundRect(left, yy - 8, max(barw, 1.4), 14, 4, fill=1, stroke=0)
        c.setFillColor(colors.HexColor(INK))
        c.setFont("Helvetica-Bold", 7)
        label = "Wav2Vec2 CTC" if i == 0 else "Whisper"
        c.drawRightString(left - 6, yy - 3, label)
        c.drawString(left + barw + 5, yy - 3, f"{value:.2f}%")
    c.setFillColor(colors.HexColor(MUTED))
    c.setFont("Helvetica", 6.5)
    c.drawCentredString((left + right) / 2, bottom - 21, "Word error rate (%)")
    return bottom - 25


def make_pdf(stats, metrics, asr_text, whisper_text):
    width, height = A4
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4, pageCompression=1)
    c.setTitle("From Voice to Text and Back - 15-page project report")
    c.setAuthor("Stutter Speech ASR and TTS project")
    c.setSubject("Speech feature extraction, ASR, CTC, WER, and TTS")
    margin = 18 * mm
    content_w = width - 2 * margin
    draw_pdf_cover(c, stats)

    # Page 2: Executive summary and objectives
    y = pdf_page(c, 2, "Executive summary", "Purpose, scope, and outcomes")
    y = pdf_para(
        c, margin, y,
        content_w,
        f"<b>Purpose.</b> This project walks through a speech technology pipeline using the current English recording "
        f"({stats['source_duration']:.1f} seconds) with reported speech disfluencies. It combines acoustic inspection, two pretrained ASR systems, a CTC "
        "decoding demonstration, word-error-rate calculation, and speech synthesis. It is an educational "
        "engineering demonstration, not a diagnostic evaluation.",
    ) - 8
    pdf_section(c, margin, y, "Activity objectives")
    y -= 12
    objectives = [
        ("01", "Inspect audio", "Visualize the waveform and log-Mel spectrum."),
        ("02", "Recognize speech", "Compare Wav2Vec2 with CTC and Whisper."),
        ("03", "Explain decoding", "Show CTC collapse/repeat and blank handling."),
        ("04", "Evaluate", "Calculate WER against a reference transcript."),
        ("05", "Synthesize", "Convert a recognized transcript back to audio."),
    ]
    cardw = (content_w - 8) / 2
    for i, (num, title, desc) in enumerate(objectives):
        xx = margin + (i % 2) * (cardw + 8)
        yy = y - (i // 2) * 50
        pdf_card(c, xx, yy, cardw, 41, f"{num}  {title}", desc)
    y -= 119
    pdf_section(c, margin, y, "End-to-end workflow")
    pdf_pipeline(c, margin, y - 5, content_w)
    pdf_para(
        c, margin, y - 57, content_w,
        "The experiment's two recognizers are parallel alternatives. Only the Wav2Vec2 path uses CTC; Whisper "
        "generates tokens with an encoder-decoder. A single recording and an unverified reference limit what can "
        "be concluded from the current scores.",
        kind="small",
    )
    c.showPage()

    # Page 3: Data and recording profile
    y = pdf_page(c, 3, "Recording and data profile", "Input characteristics and handling")
    pdf_section(c, margin, y, "Source recording")
    y -= 14
    rows = [
        ["Property", "Observed project value"],
        ["File", "audio/stutter_current.mp3"],
        ["Duration", f"{stats['source_duration']:.2f} seconds"],
        ["Decoded sample rate", f"{stats['source_rate']:,} Hz"],
        ["Analysis channels", "Mono for visual analysis and ASR"],
        ["ASR rate", "16,000 Hz for Wav2Vec2; Whisper decodes the file directly"],
        ["Primary topic", "Speaker describes a job interview and difficulty speaking by phone"],
    ]
    y = pdf_table(c, margin, y, [48 * mm, content_w - 48 * mm], rows) - 15
    y = pdf_para(
        c, margin, y, content_w,
        "<b>Preprocessing choices.</b> The plotting script loads the original sample rate and converts to mono. "
        "Wav2Vec2 uses librosa to resample to 16 kHz, its expected input rate. Whisper's audio decoder handles "
        "the MP3 and resampling internally. Keep the same source recording when comparing models.",
    ) - 9
    pdf_section(c, margin, y, "Ethical and scope notes")
    y -= 12
    y = pdf_para(
        c, margin, y, content_w,
        "Speech recordings and their transcripts can be sensitive personal data. Obtain informed permission for "
        "recording, processing, and redistribution, and share only what the speaker has authorized. The report "
        "describes observable pipeline behavior; it does not infer diagnosis, severity, or a person's identity.",
    ) - 8
    pdf_card(
        c, margin, y, content_w, 51, "Interpretation boundary",
        "A pause or repeated acoustic pattern in a plot is not, by itself, proof of a stutter. Clinical conclusions "
        "require appropriate expertise and assessment.",
        fill=PALE_MINT,
    )
    c.showPage()

    # Page 4: waveform
    y = pdf_page(c, 4, "Waveform analysis", "Amplitude plotted against time")
    pdf_section(c, margin, y, "What the waveform shows")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "The waveform represents the audio signal's amplitude over time. It helps locate broad pauses, loudness "
        "changes, and regions with speech activity. It is a useful first inspection but does not encode lexical "
        "content and cannot reliably distinguish a filled pause, a stop closure, silence, or a recording artifact "
        "without listening and further analysis.",
    ) - 8
    y = pdf_image(c, OUT / "waveform.png", margin, y, content_w, 76 * mm) - 4
    pdf_para(
        c, margin, y, content_w,
        f"Figure 1. Waveform generated from the project recording ({stats['source_duration']:.2f} s, "
        f"{stats['source_rate']:,} Hz source rate). Time is shown horizontally and normalized audio amplitude "
        "vertically.",
        "small",
    )
    c.showPage()

    # Page 5: log-Mel feature extraction
    y = pdf_page(c, 5, "Feature extraction and log-Mel spectrum", "From waveform samples to a time-frequency view")
    pdf_section(c, margin, y, "Log-Mel visualization")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "Short-time Fourier analysis decomposes local audio frames into frequency components. Mel filters combine "
        "the spectrum into perceptually spaced bands; conversion to decibels compresses the dynamic range and "
        "makes weaker components visible. This representation highlights spectral structure over time, rather "
        "than a single amplitude trace.",
    ) - 6
    y = pdf_image(c, OUT / "log_mel_spectrogram.png", margin, y, content_w, 74 * mm) - 3
    pdf_para(
        c, margin, y, content_w,
        "Figure 2. Log-Mel power spectrogram saved by main.py. Brighter regions indicate stronger relative "
        "spectral energy; frequency is displayed on the Mel scale.",
        "small",
    )
    y -= 18
    pdf_section(c, margin, y, "Feature parameters")
    y -= 12
    pdf_table(
        c, margin, y,
        [42 * mm, 31 * mm, content_w - 73 * mm],
        [
            ["Parameter", "Value", "Role"],
            ["Mel bands", "128", "Frequency bins in perceptual scale"],
            ["FFT window", "2048 samples", "Local spectral analysis window"],
            ["Hop length", "512 samples", "Frame advance and time resolution"],
            ["Reference level", "Maximum power", "Relative decibel scaling"],
        ],
    )
    c.showPage()

    # Page 6: model architecture
    y = pdf_page(c, 6, "ASR architecture", "Two recognizer branches, one evaluation/synthesis path")
    pdf_section(c, margin, y, "Pipeline architecture")
    y -= 11
    pdf_pipeline(c, margin, y, content_w)
    y -= 52
    branch_rows = [
        ["Branch", "Input representation", "Recognition / decoder"],
        ["Wav2Vec2", "16 kHz mono waveform", "Context encoder > character logits > CTC"],
        ["Whisper", "Decoded audio / log-Mel features", "Encoder > autoregressive text decoder"],
        ["Shared evaluation", "Reference + hypothesis text", "Normalized words > edit counts > WER"],
        ["TTS path", "Selected transcript", "Text synthesizer > generated speech waveform"],
    ]
    y = pdf_table(c, margin, y, [31 * mm, 58 * mm, content_w - 89 * mm], branch_rows) - 12
    pdf_card(
        c, margin, y, content_w, 54, "Keep the model distinction clear",
        "CTC is the Wav2Vec2 decoding method demonstrated here. Whisper is an encoder-decoder system and does not "
        "pass through the project's CTC decoder. They are not sequential components of one recognizer.",
    )
    c.showPage()

    # Page 7: Wav2Vec2
    y = pdf_page(c, 7, "Wav2Vec2 ASR", "Self-supervised speech representations with CTC output decoding")
    pdf_section(c, margin, y, "Model and inference")
    y -= 12
    y = pdf_para(
        c, margin, y, content_w,
        "<b>Checkpoint:</b> facebook/wav2vec2-base-960h. The project uses the matching Wav2Vec2 processor and "
        "ForCTC model. Mono audio is resampled to 16 kHz and passed as input_values. The model returns frame-wise "
        "logits; argmax selects the most likely token per frame; processor.batch_decode applies CTC collapsing "
        "and converts tokens to text.",
    ) - 8
    wav_excerpt = asr_text[:390] + ("..." if len(asr_text) > 390 else "")
    y = pdf_card(
        c, margin, y, content_w, 87,
        "Recognized output (excerpt)",
        clean_markup(wav_excerpt),
        fill=PALE,
    )
    y -= 14
    pdf_section(c, margin, y, "Observation")
    y -= 12
    pdf_para(
        c, margin, y, content_w,
        "The current result includes substitutions and inserted/repeated-looking words. That is plausible in "
        "spontaneous and disfluent speech, but each suspected error must be compared to the actual audio. A "
        "recognizer may also normalize or omit repetitions, so its output is not a verbatim ground truth.",
    )
    c.showPage()

    # Page 8: CTC
    y = pdf_page(c, 8, "CTC decoding explained", "Connectionist Temporal Classification and the blank token")
    pdf_section(c, margin, y, "Why CTC is useful")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "During training/inference, an acoustic encoder produces a sequence of frame-level token distributions "
        "whose length is much greater than the transcript. CTC sums over valid alignments that collapse to the "
        "same target, avoiding the need to label an exact time for every character. Greedy decoding is a simple "
        "approximation that selects the best token independently at each frame.",
    ) - 15
    pdf_section(c, margin, y, "Greedy collapse example")
    y -= 16
    steps = [
        ("Frame path", "blank | A | A | blank | A | B | B | blank"),
        ("Merge adjacent duplicates", "blank | A | blank | A | B | blank"),
        ("Remove blank symbols", "A | A | B"),
        ("Decoded string", "AAB"),
    ]
    row_y = y
    for idx, (label, value) in enumerate(steps):
        c.setFillColor(colors.HexColor(PALE if idx % 2 == 0 else PALE_MINT))
        c.roundRect(margin, row_y - 22, content_w, 22, 4, fill=1, stroke=0)
        pdf_plain(c, margin + 8, row_y - 14, label, 8, BLUE, True)
        pdf_plain(c, margin + 62 * mm, row_y - 14, value, 8, INK)
        row_y -= 29
    y = row_y - 3
    pdf_card(
        c, margin, y, content_w, 62, "Why the blank matters",
        "A A with no blank collapses to A. A blank separates repeated labels: A, blank, A becomes AA. "
        "The correct order is to merge adjacent repeats first, then remove blanks.",
    )
    c.showPage()

    # Page 9: Whisper
    y = pdf_page(c, 9, "Whisper ASR", "Encoder-decoder transcription with token generation")
    pdf_section(c, margin, y, "Inference setup")
    y -= 12
    y = pdf_para(
        c, margin, y, content_w,
        "<b>Implementation:</b> faster-whisper, base model, CPU device, int8 compute, English language hint, and "
        "beam size 5. Whisper's audio processing creates log-Mel features internally and its decoder generates "
        "tokens conditioned on the encoded audio and previous tokens. It is architecturally different from "
        "Wav2Vec2-CTC.",
    ) - 8
    excerpt = whisper_text[:450] + ("..." if len(whisper_text) > 450 else "")
    y = pdf_card(
        c, margin, y, content_w, 100,
        "Recognized output (excerpt)",
        clean_markup(excerpt),
        fill=PALE_MINT,
    )
    y -= 14
    pdf_section(c, margin, y, "What this output represents")
    y -= 12
    pdf_para(
        c, margin, y, content_w,
        "This is a model-generated hypothesis, not a verified transcript. Fluent-looking punctuation and wording "
        "can conceal missed disfluencies or recognition errors. Always compare the text with the waveform and "
        "recording, especially when evaluating speech patterns.",
    )
    c.showPage()

    # Page 10: WER methodology
    y = pdf_page(c, 10, "WER methodology", "Quantifying word-level transcription differences")
    pdf_section(c, margin, y, "Definition")
    y -= 14
    y = pdf_card(
        c, margin, y, content_w, 54, "Word error rate",
        "WER = (S + D + I) / N. S = substitutions; D = deletions; I = insertions; N = number of reference words.",
        fill=PALE_MINT,
    ) - 17
    pdf_section(c, margin, y, "How the project computes it")
    y -= 12
    rows = [
        ["Step", "Operation"],
        ["1. Reference", "Load outputs/reference_transcript.txt (currently copied from Whisper)."],
        ["2. Hypotheses", "Read Wav2Vec2 and Whisper transcript output files."],
        ["3. Normalize", "Lowercase and remove punctuation consistently."],
        ["4. Align", "JiWER aligns token sequences and counts edit operations."],
        ["5. Report", "Print S/D/I and WER; save outputs/wer_results.csv."],
    ]
    y = pdf_table(c, margin, y, [36 * mm, content_w - 36 * mm], rows) - 17
    pdf_section(c, margin, y, "Interpretation")
    y -= 12
    pdf_para(
        c, margin, y, content_w,
        "WER is sensitive to the reference transcript and tokenization conventions. A value of 0.0 means exact "
        "normalized word agreement on this reference; a value above 1.0 is possible when insertions are numerous. "
        "Punctuation is typically excluded, but decisions about repetitions, partial words, and filled pauses "
        "must be consistent and documented.",
    )
    c.showPage()

    # Page 11: WER results
    y = pdf_page(c, 11, "ASR comparison results", "Observed scores on this one recording and provisional reference")
    pdf_section(c, margin, y, "Current result table")
    y -= 14
    table_rows = [["Model", "Hits", "Sub.", "Del.", "Ins.", "Ref words", "WER"]]
    for r in metrics:
        table_rows.append([
            r["model"], r["hits"], r["substitutions"], r["deletions"],
            r["insertions"], r["reference_words"], f"{float(r['wer']):.2%}",
        ])
    y = pdf_table(
        c, margin, y,
        [41 * mm, 18 * mm, 18 * mm, 18 * mm, 18 * mm, 28 * mm, content_w - 141 * mm],
        table_rows,
    ) - 17
    pdf_section(c, margin, y, "WER visual comparison")
    y -= 5
    y = pdf_bar_chart(c, margin, y, content_w, metrics) - 3
    y = pdf_para(
        c, margin, y, content_w,
        f"On this provisional {metrics[0]['reference_words']}-word Whisper-derived reference, Wav2Vec2 differs by "
        f"{metrics[0]['substitutions']} substitutions, {metrics[0]['deletions']} deletions, and "
        f"{metrics[0]['insertions']} insertions ({float(metrics[0]['wer']):.2%} WER). Whisper scores "
        f"{float(metrics[1]['wer']):.2%} because its transcript is the reference. These are agreement scores, "
        "not verified accuracy measurements.",
    ) - 9
    pdf_card(
        c, margin, y, content_w, 66, "Critical caveat",
        "The reference is copied directly from Whisper and has not been independently verified. Whisper therefore "
        "scores zero by construction. Wav2Vec2's score measures disagreement with Whisper, not its true error rate. "
        "Do not interpret these figures as model accuracy.",
        fill="#FFF4D6",
    )
    c.showPage()

    # Page 12: qualitative analysis
    y = pdf_page(c, 12, "Qualitative output analysis", "Differences between recognizer hypotheses")
    pdf_section(c, margin, y, "Side-by-side excerpts")
    y -= 13
    w_excerpt = whisper_text[:330] + ("..." if len(whisper_text) > 330 else "")
    a_excerpt = asr_text[:330] + ("..." if len(asr_text) > 330 else "")
    half = (content_w - 8) / 2
    pdf_card(c, margin, y, half, 153, "Wav2Vec2 hypothesis", clean_markup(a_excerpt), fill=PALE)
    pdf_card(c, margin + half + 8, y, half, 153, "Whisper hypothesis", clean_markup(w_excerpt), fill=PALE_MINT)
    y -= 170
    pdf_section(c, margin, y, "What the current comparison can and cannot show")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "The output texts differ in recognition choices and in how fluency and punctuation are rendered. Such "
        "differences can reflect model design, training data, decoding, and normalization. The current WER "
        "calculation compares each output to the Whisper transcript itself; it does not establish which "
        "words were actually spoken. A defensible analysis requires an independently prepared reference and "
        "manual adjudication of ambiguous segments.",
    ) - 9
    pdf_card(
        c, margin, y, content_w, 65, "Recommended transcript protocol",
        "Listen to the source audio; transcribe verbatim; define a consistent convention for repetitions, partial "
        "words, and fillers; have a second reviewer resolve uncertain spans; then calculate WER for both systems.",
    )
    c.showPage()

    # Page 13: TTS
    y = pdf_page(c, 13, "Text-to-speech output", "Completing the return path from text to audio")
    pdf_section(c, margin, y, "Synthesis procedure")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "The PowerShell script reads outputs/whisper_transcription.txt and uses Windows System.Speech to synthesize "
        "the recognized text. The generated audio is saved as outputs/tts_output.wav. The notebook also offers "
        "an optional gTTS path, which requires network access.",
    ) - 10
    rows = [
        ["Artifact", "Observed properties"],
        ["Generated file", "outputs/tts_output.wav"],
        ["Encoding", "PCM WAV"],
        ["Sample rate", f"{stats['tts_rate']:,} Hz"],
        ["Channels", str(stats["tts_channels"])],
        ["Duration", f"{stats['tts_duration']:.2f} seconds"],
        ["Text source", "Whisper transcript"],
    ]
    y = pdf_table(c, margin, y, [44 * mm, content_w - 44 * mm], rows) - 17
    pdf_section(c, margin, y, "Interpretation")
    y -= 13
    pdf_card(
        c, margin, y, content_w, 79, "Synthetic voice is not the source speaker",
        "TTS demonstrates text-to-speech conversion only. It does not reproduce the source speaker's vocal identity, "
        "prosody, or stuttering. The output is a separate synthesized performance of the transcript.",
        fill=PALE_MINT,
    )
    c.showPage()

    # Page 14: reproducibility
    y = pdf_page(c, 14, "Reproducibility and project structure", "How to rerun the pipeline")
    pdf_section(c, margin, y, "Core files")
    y -= 13
    files = [
        ["File", "Purpose"],
        ["main.py", "Waveform and log-Mel figures"],
        ["asr.py", "Wav2Vec2 CTC transcription"],
        ["whisper_asr.py", "Whisper transcription"],
        ["ctc_demo.py", "Illustrative greedy CTC collapse"],
        ["wer.py", "Two-model WER and CSV export"],
        ["tts.ps1", "Windows speech synthesis"],
        ["stutter_speech_pipeline.ipynb", "Notebook/Colab workflow"],
        ["report/generate_pdf.py", "Two-page mini-report generator"],
    ]
    y = pdf_table(c, margin, y, [70 * mm, content_w - 70 * mm], files) - 15
    pdf_section(c, margin, y, "Run order (PowerShell)")
    y -= 13
    code = [
        'Set-Location "D:\\stutter project 191"',
        '$py = ".\\.venv\\Scripts\\python.exe"',
        '& $py main.py',
        '& $py asr.py',
        '& $py whisper_asr.py',
        '& $py wer.py',
        ".\\tts.ps1",
    ]
    c.setFillColor(colors.HexColor(NAVY))
    c.roundRect(margin, y - 73, content_w, 73, 5, fill=1, stroke=0)
    for i, line in enumerate(code):
        c.setFillColor(colors.HexColor("#D8E8F1"))
        c.setFont("Courier", 7.4)
        c.drawString(margin + 10, y - 14 - i * 9, line)
    y -= 88
    pdf_para(
        c, margin, y, content_w,
        "Model checkpoints may download on first run. WER results should be regenerated only after verifying the "
        "reference file against the source recording.",
        "small",
    )
    c.showPage()

    # Page 15: conclusions and references
    y = pdf_page(c, 15, "Conclusions, limitations, and next steps", "Responsible interpretation and suggested improvements")
    pdf_section(c, margin, y, "Conclusions")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "The project demonstrates acoustic visualization, two contrasting ASR approaches, CTC decoding, WER "
        "calculation, and TTS. Its outputs make the pipeline concrete and reproducible. The current run shows "
        "different hypotheses from Wav2Vec2 and Whisper, but current WER values are not independently validated.",
    ) - 7
    pdf_section(c, margin, y, "Limitations")
    y -= 13
    limitations = [
        f"Single {stats['source_duration']:.1f}-second recording; results cannot be generalized to other speakers or conditions.",
        "Reference transcript is copied from Whisper and has not been independently verified.",
        "ASR may normalize or omit disfluencies; transcript conventions affect WER.",
        "Greedy CTC does not use a language-model beam search in this implementation.",
        "TTS output is synthetic and not a reconstruction of the original voice.",
    ]
    for item in limitations:
        y = pdf_para(c, margin + 4, y, content_w - 4, f"- {clean_markup(item)}", "body") - 2
    y -= 5
    pdf_section(c, margin, y, "Priority next steps")
    y -= 13
    y = pdf_para(
        c, margin, y, content_w,
        "1. Create and independently review a verbatim reference transcript.<br/>"
        "2. Define a consistent disfluency and punctuation normalization policy.<br/>"
        "3. Re-run WER and report edit counts with the verified reference.<br/>"
        "4. Expand evaluation to more recordings and report per-speaker and aggregate results.<br/>"
        "5. Preserve consent, privacy, and clear separation between technical outputs and clinical claims.",
    ) - 6
    pdf_section(c, margin, y, "Selected references")
    y -= 12
    refs = [
        "Graves, A. et al. (2006). Connectionist Temporal Classification: Labelling Unsegmented Sequence Data with Recurrent Neural Networks. ICML.",
        "Baevski, A. et al. (2020). wav2vec 2.0: A Framework for Self-Supervised Learning of Speech Representations. NeurIPS.",
        "Radford, A. et al. (2023). Robust Speech Recognition via Large-Scale Weak Supervision. ICML.",
        "JiWER documentation. Word error rate and alignment measures for automatic speech recognition evaluation.",
    ]
    for ref in refs:
        y = pdf_para(c, margin + 3, y, content_w - 3, clean_markup(ref), "small") - 3
    c.showPage()
    c.save()


def ppt_rgb(hex_color):
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def ppt_shape(slide, kind, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_rgb(fill)
    shape.line.color.rgb = ppt_rgb(line or fill)
    if radius and shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = 0.08
    return shape


def ppt_text(slide, x, y, w, h, text, size=16, color=INK, bold=False, font="Aptos",
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = ppt_rgb(color)
        p.space_after = Pt(5)
    return box


def ppt_add_title(slide, title, subtitle, n):
    ppt_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.12, CYAN, CYAN, radius=False)
    ppt_text(slide, 0.58, 0.31, 11.8, 0.52, title, size=25, color=NAVY, bold=True)
    ppt_text(slide, 0.61, 0.88, 12.0, 0.35, subtitle, size=11, color=MUTED)
    ppt_shape(slide, MSO_SHAPE.RECTANGLE, 0.6, 7.14, 12.15, 0.012, LIGHT_LINE, LIGHT_LINE, radius=False)
    ppt_text(slide, 0.62, 7.2, 9.8, 0.18, "STUTTER SPEECH ASR + TTS | Educational demonstration", size=7.5, color=MUTED)
    ppt_text(slide, 11.9, 7.18, 0.8, 0.2, f"{n:02d} / 15", size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def ppt_new_slide(prs, title, subtitle, n):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ppt_rgb(WHITE)
    ppt_add_title(slide, title, subtitle, n)
    return slide


def ppt_card(slide, x, y, w, h, title, body, accent=BLUE, body_size=13):
    ppt_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, PALE, LIGHT_LINE)
    ppt_shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.12, 0.07, h - 0.24, accent, accent, radius=False)
    ppt_text(slide, x + 0.22, y + 0.16, w - 0.42, 0.42, title, size=15, color=NAVY, bold=True)
    ppt_text(slide, x + 0.22, y + 0.66, w - 0.42, h - 0.78, body, size=body_size, color=INK)


def ppt_image(slide, path, x, y, w, h):
    from PIL import Image as PILImage

    with PILImage.open(path) as im:
        iw, ih = im.size
    ratio = min(w / iw, h / ih)
    dw, dh = iw * ratio, ih * ratio
    return slide.shapes.add_picture(str(path), Inches(x + (w - dw) / 2), Inches(y + (h - dh) / 2), width=Inches(dw), height=Inches(dh))


def ppt_flow(slide, y=2.1, labels=None):
    labels = labels or [
        ("Speech", "MP3"),
        ("Features", "Waveform + Mel"),
        ("ASR", "2 alternatives"),
        ("Text", "CTC / WER"),
        ("TTS", "Synthesized audio"),
    ]
    x0 = 0.65
    gap = 0.34
    w = 2.22
    for i, (title, sub) in enumerate(labels):
        x = x0 + i * (w + gap)
        ppt_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 1.0, PALE if i % 2 == 0 else PALE_MINT, BLUE if i % 2 == 0 else CYAN)
        ppt_text(slide, x + 0.1, y + 0.2, w - 0.2, 0.3, title, 14, NAVY, True, align=PP_ALIGN.CENTER)
        ppt_text(slide, x + 0.1, y + 0.57, w - 0.2, 0.24, sub, 9, MUTED, align=PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            ppt_text(slide, x + w, y + 0.33, gap, 0.25, ">", 16, BLUE, True, align=PP_ALIGN.CENTER)


def ppt_chart(slide, rows, x=2.1, y=2.3, w=9.3):
    left = x
    chartw = w
    for tick in (0, 20, 40, 60, 80, 100):
        xx = left + chartw * tick / 100
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xx), Inches(y + 0.42), Inches(xx), Inches(y + 2.1))
        line.line.color.rgb = ppt_rgb(LIGHT_LINE)
        ppt_text(slide, xx - 0.15, y + 2.12, 0.3, 0.2, str(tick), 9, MUTED, align=PP_ALIGN.CENTER)
    for i, row in enumerate(rows[:2]):
        label = "Wav2Vec2 CTC" if row["model"].startswith("Wav2Vec2") else "Whisper"
        value = float(row["wer"]) * 100
        yy = y + 0.65 + i * 0.72
        ppt_text(slide, 0.65, yy, 1.35, 0.32, label, 12, INK, True)
        bw = max(chartw * value / 100, 0.08)
        ppt_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, yy - 0.02, bw, 0.4, BLUE if i == 0 else CORAL, BLUE if i == 0 else CORAL)
        ppt_text(slide, left + bw + 0.12, yy + 0.03, 1.2, 0.28, f"{value:.2f}%", 12, NAVY, True)
    ppt_text(slide, x, y + 2.4, w, 0.22, "WER (%) - lower is better only when the reference is trustworthy", 9, MUTED, align=PP_ALIGN.CENTER)


def ppt_table(slide, rows, x, y, w, h, font_size=12):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    for col in range(len(rows[0])):
        table.columns[col].width = Inches(w / len(rows[0]))
    for r, row in enumerate(rows):
        for col, value in enumerate(row):
            cell = table.cell(r, col)
            cell.text = str(value)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ppt_rgb(NAVY if r == 0 else (WHITE if r % 2 else PALE))
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size if r else font_size - 1)
                p.font.bold = r == 0
                p.font.color.rgb = ppt_rgb(WHITE if r == 0 else INK)
    return table_shape


def make_ppt(stats, metrics, asr_text, whisper_text):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "From Voice to Text and Back: ASR, CTC, WER and TTS"
    prs.core_properties.subject = "Professional project presentation"
    prs.core_properties.author = "Stutter Speech ASR and TTS project"

    # Slide 1: cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ppt_rgb(NAVY)
    ppt_shape(slide, MSO_SHAPE.OVAL, 9.45, -1.3, 5.0, 5.0, BLUE, BLUE, radius=False)
    ppt_shape(slide, MSO_SHAPE.OVAL, 10.4, -0.5, 3.2, 3.2, CYAN, CYAN, radius=False)
    ppt_text(slide, 0.9, 0.75, 8.0, 0.35, "SPEECH TECHNOLOGY PROJECT", 12, MINT, True)
    ppt_text(slide, 0.9, 1.55, 10.3, 1.35, "From Voice to Text\nand Back", 34, WHITE, True)
    ppt_text(slide, 0.95, 3.0, 8.7, 0.7, "ASR, CTC, WER and text-to-speech analysis\nof a speech sample with disfluencies", 18, "#D8E8F1")
    ppt_flow(slide, y=4.35)
    ppt_text(slide, 0.95, 5.85, 8.7, 0.55, f"140.35-second English sample  |  Wav2Vec2 + Whisper  |  TTS", 12, WHITE)
    ppt_text(slide, 0.95, 6.72, 9.0, 0.25, "Educational demonstration - not a clinical assessment", 9, "#D8E8F1")
    ppt_text(slide, 12.0, 7.16, 0.8, 0.2, "01 / 15", 8, "#D8E8F1", align=PP_ALIGN.RIGHT)

    # Slide 2
    s = ppt_new_slide(prs, "The project in one minute", "A complete speech-to-text-to-speech workflow", 2)
    ppt_flow(s, 1.75)
    cards = [
        ("Analyze", "Inspect amplitude and frequency structure.", BLUE),
        ("Recognize", "Compare CTC-based Wav2Vec2 with Whisper.", CYAN),
        ("Evaluate", "Compute edit counts and word error rate.", MINT),
        ("Synthesize", "Speak the selected transcript back as audio.", CORAL),
    ]
    for i, (t, b, a) in enumerate(cards):
        ppt_card(s, 0.7 + i * 3.15, 3.55, 2.9, 1.8, t, b, a, 12)
    ppt_text(s, 0.78, 5.8, 11.8, 0.75, f"Provisional agreement vs. Whisper reference: Wav2Vec2 {float(metrics[0]['wer']):.2%}; Whisper {float(metrics[1]['wer']):.2%}. These are not accuracy scores.", 15, NAVY, True)

    # Slide 3
    s = ppt_new_slide(prs, "Input recording and experiment scope", "What was analyzed - and what the numbers represent", 3)
    ppt_card(s, 0.7, 1.55, 3.7, 1.5, "Recording", f"audio/stutter_current.mp3\n{stats['source_duration']:.2f} seconds | {stats['source_rate']:,} Hz", BLUE, 13)
    ppt_card(s, 4.8, 1.55, 3.7, 1.5, "Processing", "Mono analysis\nWav2Vec2 resampled to 16 kHz", CYAN, 13)
    ppt_card(s, 8.9, 1.55, 3.7, 1.5, "Scope", "One English speech sample\nEducational pipeline demonstration", MINT, 13)
    ppt_text(s, 0.8, 3.65, 11.7, 0.55, "Research question", 16, BLUE, True)
    ppt_text(s, 0.8, 4.15, 11.2, 0.85, "How do two popular ASR approaches transcribe the same disfluent speech sample, and how can their word-level differences be measured?", 22, NAVY, True)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 5.48, 11.7, 0.9, "#FFF4D6", "#D6A94D")
    ppt_text(s, 1.0, 5.65, 11.3, 0.55, "This single-sample exercise cannot diagnose a speaker or generalize model accuracy. WER depends on a verified reference transcript.", 13, INK)

    # Slide 4 waveform
    s = ppt_new_slide(prs, "1 | Inspect the waveform", "Amplitude and pauses over the recording timeline", 4)
    ppt_image(s, OUT / "waveform.png", 0.82, 1.5, 11.7, 4.5)
    ppt_text(s, 0.9, 6.18, 11.6, 0.55, "Waveform is useful for locating activity and pauses. It does not reveal the words or diagnose a fluency disorder by itself.", 13, INK)

    # Slide 5 spectrogram
    s = ppt_new_slide(prs, "2 | Transform audio into log-Mel features", "Short-time spectral energy represented on a perceptual frequency scale", 5)
    ppt_image(s, OUT / "log_mel_spectrogram.png", 0.75, 1.42, 11.8, 4.35)
    ppt_card(s, 0.9, 5.85, 3.65, 0.8, "Mel bands", "128", BLUE, 11)
    ppt_card(s, 4.85, 5.85, 3.65, 0.8, "FFT / hop", "2048 / 512 samples", CYAN, 11)
    ppt_card(s, 8.8, 5.85, 3.65, 0.8, "Scale", "Power converted to dB", MINT, 11)

    # Slide 6 architecture
    s = ppt_new_slide(prs, "3 | Two ASR architectures", "Parallel alternatives; not a single serial recognizer", 6)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 3.0, 1.8, 1.0, PALE, BLUE)
    ppt_text(s, 0.82, 3.2, 1.55, 0.55, "Speech\nMP3", 14, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.45, 1.65, 3.25, 1.25, PALE_MINT, CYAN)
    ppt_text(s, 3.6, 1.88, 2.95, 0.75, "Wav2Vec2 encoder\nwaveform -> character logits", 15, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.1, 1.65, 2.25, 1.25, "#FFF4D6", GOLD)
    ppt_text(s, 7.25, 1.88, 1.95, 0.75, "CTC collapse\nrepeat + blank", 14, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 3.45, 4.2, 3.25, 1.25, PALE, BLUE)
    ppt_text(s, 3.6, 4.43, 2.95, 0.75, "Whisper encoder\nlog-Mel audio features", 15, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.1, 4.2, 2.25, 1.25, PALE_MINT, MINT)
    ppt_text(s, 7.25, 4.43, 1.95, 0.75, "Text decoder\nautoregressive", 14, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 10.15, 3.0, 2.3, 1.0, PALE, CORAL)
    ppt_text(s, 10.3, 3.22, 2.0, 0.55, "Text -> WER -> TTS", 14, NAVY, True, align=PP_ALIGN.CENTER)
    for x1, y1, x2, y2 in [
        (2.5, 3.35, 3.45, 2.28), (2.5, 3.55, 3.45, 4.8),
        (6.7, 2.28, 7.1, 2.28), (9.35, 2.28, 10.15, 3.25),
        (6.7, 4.8, 7.1, 4.8), (9.35, 4.8, 10.15, 3.75),
    ]:
        connector = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        connector.line.color.rgb = ppt_rgb(BLUE)
        connector.line.width = Pt(2)
    ppt_text(s, 0.82, 6.0, 11.6, 0.45, "Important: Whisper is not a CTC model. The Wav2Vec2 branch is the one decoded with CTC.", 14, CORAL, True)

    # Slide 7 Wav2Vec2
    s = ppt_new_slide(prs, "4 | Wav2Vec2 with CTC", "facebook/wav2vec2-base-960h - greedy token decoding", 7)
    ppt_card(s, 0.8, 1.55, 3.6, 1.45, "Input", "Mono waveform resampled to 16 kHz", BLUE, 14)
    ppt_card(s, 4.85, 1.55, 3.6, 1.45, "Acoustic model", "Convolutional features + Transformer context", CYAN, 14)
    ppt_card(s, 8.9, 1.55, 3.6, 1.45, "Decoder", "Frame-wise argmax + CTC batch decode", MINT, 14)
    excerpt = asr_text[:340] + ("..." if len(asr_text) > 340 else "")
    ppt_card(s, 0.8, 3.5, 11.7, 2.1, "Transcript excerpt", excerpt, BLUE, 14)
    ppt_text(s, 0.9, 5.9, 11.7, 0.6, "Observed transcription contains substitutions and extra tokens; confirm errors against audio rather than assuming the reference is correct.", 13, MUTED)

    # Slide 8 CTC
    s = ppt_new_slide(prs, "5 | CTC: align without frame-level labels", "Collapse adjacent repeats first, then remove blank tokens", 8)
    steps = [
        ("Frame predictions", "blank | A | A | blank | A | B | B | blank", BLUE),
        ("Merge adjacent repeats", "blank | A | blank | A | B | blank", CYAN),
        ("Remove blanks", "A | A | B", MINT),
        ("Decoded output", "AAB", CORAL),
    ]
    for i, (head, body, col) in enumerate(steps):
        yy = 1.55 + i * 1.05
        ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, yy, 11.4, 0.78, PALE if i % 2 == 0 else PALE_MINT, col)
        ppt_text(s, 1.18, yy + 0.1, 2.5, 0.52, head, 13, NAVY, True)
        ppt_text(s, 3.75, yy + 0.12, 8.2, 0.5, body, 16, INK, font="Consolas")
    ppt_text(s, 1.0, 6.05, 11.4, 0.46, "A, blank, A -> AA; adjacent A, A -> A. The blank preserves repeated letters when separated.", 14, NAVY, True)

    # Slide 9 Whisper
    s = ppt_new_slide(prs, "6 | Whisper-base", "Encoder-decoder ASR with autoregressive text token generation", 9)
    ppt_card(s, 0.85, 1.55, 3.55, 1.35, "Audio features", "Whisper extracts log-Mel features internally", BLUE, 13)
    ppt_card(s, 4.9, 1.55, 3.55, 1.35, "Model", "faster-whisper base; CPU int8", CYAN, 13)
    ppt_card(s, 8.95, 1.55, 3.55, 1.35, "Decoding", "English hint; beam size 5", MINT, 13)
    excerpt = whisper_text[:390] + ("..." if len(whisper_text) > 390 else "")
    ppt_card(s, 0.85, 3.35, 11.65, 2.3, "Transcript excerpt", excerpt, MINT, 15)
    ppt_text(s, 0.95, 5.9, 11.4, 0.58, "Fluent wording is not proof of verbatim transcription. The recognizer may normalize or omit disfluencies.", 13, MUTED)

    # Slide 10 WER methodology
    s = ppt_new_slide(prs, "7 | Evaluation with word error rate", "Word-level edit distance relative to a reference transcript", 10)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 1.55, 11.55, 1.05, PALE_MINT, CYAN)
    ppt_text(s, 1.25, 1.77, 10.8, 0.6, "WER = (Substitutions + Deletions + Insertions) / Reference words", 23, NAVY, True, align=PP_ALIGN.CENTER)
    stages = [
        ("Reference", f"{metrics[0]['reference_words']} words"),
        ("Normalize", "case / punctuation"),
        ("Align", "JiWER operations"),
        ("Count", "S, D, I"),
    ]
    for i, (a, b) in enumerate(stages):
        x = 0.95 + i * 3.1
        ppt_card(s, x, 3.25, 2.7, 1.3, a, b, [BLUE, CYAN, MINT, CORAL][i], 12)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.15, 11.55, 0.95, "#FFF4D6", "#D6A94D")
    ppt_text(s, 1.1, 5.35, 11.1, 0.52, "Current reference is a Whisper copy and is unverified; treat the scores as agreement, not accuracy.", 14, INK, True)

    # Slide 11 results
    s = ppt_new_slide(prs, "8 | Results on the current reference", "One sample; current reference transcript is provisional", 11)
    ppt_chart(s, metrics)
    rows = [["Model", "Hits", "Sub", "Del", "Ins", "Words", "WER"]]
    for r in metrics:
        rows.append([r["model"], r["hits"], r["substitutions"], r["deletions"], r["insertions"], r["reference_words"], f"{float(r['wer']):.2%}"])
    ppt_table(s, rows, 0.8, 5.15, 11.7, 1.2, font_size=11)

    # Slide 12 caveat
    s = ppt_new_slide(prs, "9 | Critical evaluation caveat", "Whisper transcript is used as the provisional reference", 12)
    ppt_card(s, 0.8, 1.55, 5.7, 1.55, "Current reference", f"{metrics[0]['reference_words']} words; copied from Whisper, not independently verified", GOLD, 14)
    ppt_card(s, 6.8, 1.55, 5.7, 1.55, "Whisper hypothesis", "Matches the reference by construction", CORAL, 14)
    ppt_text(s, 0.95, 3.55, 11.4, 0.6, f"Whisper's {float(metrics[1]['wer']):.2%} is not a fair accuracy benchmark.", 24, CORAL, True, align=PP_ALIGN.CENTER)
    ppt_text(s, 1.1, 4.38, 11.1, 1.2, f"Wav2Vec2's {float(metrics[0]['wer']):.2%} measures disagreement with Whisper, not verified error. Prepare a verbatim transcript independently, then recalculate both scores.", 18, NAVY, align=PP_ALIGN.CENTER)
    ppt_text(s, 1.0, 6.15, 11.4, 0.35, "These results describe one example only; do not infer general model superiority.", 12, MUTED, align=PP_ALIGN.CENTER)

    # Slide 13 TTS
    s = ppt_new_slide(prs, "10 | Text-to-speech", "Convert a recognized transcript back into an audio signal", 13)
    ppt_flow(s, 1.9, [
        ("Whisper text", "recognized words"),
        ("TTS engine", "Windows SAPI"),
        ("WAV output", "PCM audio"),
        ("Listen", "quality check"),
        ("Caution", "synthetic voice"),
    ])
    ppt_card(s, 0.9, 3.55, 5.5, 1.65, "Generated artifact", f"outputs/tts_output.wav\n{stats['tts_duration']:.1f} seconds | {stats['tts_rate']:,} Hz | mono", BLUE, 14)
    ppt_card(s, 6.85, 3.55, 5.5, 1.65, "Interpretation", "Synthetic speech of the transcript; it does not recreate the original speaker or their disfluencies.", MINT, 14)
    ppt_text(s, 1.0, 5.8, 11.4, 0.55, "The notebook includes an optional gTTS alternative; it requires an internet connection.", 13, MUTED)

    # Slide 14 reproducibility
    s = ppt_new_slide(prs, "11 | Reproducible run order", "Repository scripts and generated artifacts", 14)
    ppt_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 1.45, 5.7, 4.8, NAVY, NAVY)
    commands = [
        'Set-Location "D:\\stutter project 191"',
        '$py = ".\\.venv\\Scripts\\python.exe"',
        '& $py main.py',
        '& $py asr.py',
        '& $py whisper_asr.py',
        '& $py wer.py',
        ".\\tts.ps1",
    ]
    for i, command in enumerate(commands):
        ppt_text(s, 1.05, 1.78 + i * 0.56, 5.2, 0.35, command, 11, "#D8E8F1", font="Consolas")
    ppt_card(s, 6.9, 1.45, 5.55, 1.4, "Scripts", "main.py | asr.py | whisper_asr.py | wer.py | tts.ps1", BLUE, 12)
    ppt_card(s, 6.9, 3.1, 5.55, 1.4, "Artifacts", "Waveform, log-Mel image, transcripts, WER CSV, synthesized WAV", CYAN, 12)
    ppt_card(s, 6.9, 4.75, 5.55, 1.4, "Reproducibility", "Use same audio and a human-verified reference; record model versions and settings.", MINT, 12)

    # Slide 15 conclusion
    s = ppt_new_slide(prs, "Takeaways and next steps", "A useful pipeline demonstration - with evaluation caveats made explicit", 15)
    ppt_card(s, 0.8, 1.55, 3.7, 1.55, "What works", "Feature plots, two ASR paths, CTC explanation, provisional comparison, and TTS.", BLUE, 13)
    ppt_card(s, 4.8, 1.55, 3.7, 1.55, "What needs validation", "Reference transcript and the resulting accuracy comparison.", CORAL, 13)
    ppt_card(s, 8.8, 1.55, 3.7, 1.55, "What comes next", "Independent transcript review and broader recordings.", MINT, 13)
    ppt_text(s, 1.05, 3.7, 11.2, 0.6, "Bottom line", 16, BLUE, True, align=PP_ALIGN.CENTER)
    ppt_text(s, 1.2, 4.25, 10.9, 1.0, "A complete speech pipeline is demonstrated. Reliable ASR conclusions require an independent, carefully annotated reference and evaluation beyond one sample.", 22, NAVY, True, align=PP_ALIGN.CENTER)
    ppt_text(s, 1.2, 5.75, 10.9, 0.75, "Selected reading: Graves et al. (2006), CTC; Baevski et al. (2020), wav2vec 2.0; Radford et al. (2023), Whisper.", 11, MUTED, align=PP_ALIGN.CENTER)

    prs.save(PPT_PATH)


def main():
    REPORT.mkdir(parents=True, exist_ok=True)
    stats = audio_stats()
    metrics = load_metrics()
    asr_text = read_transcript("asr_transcription.txt")
    whisper_text = read_transcript("whisper_transcription.txt")
    if not asr_text or not whisper_text:
        raise FileNotFoundError("Run both ASR scripts before generating the presentation and report.")
    make_pdf(stats, metrics, asr_text, whisper_text)
    make_ppt(stats, metrics, asr_text, whisper_text)
    print(f"PDF report: {PDF_PATH}")
    print(f"PowerPoint deck: {PPT_PATH}")


if __name__ == "__main__":
    main()
